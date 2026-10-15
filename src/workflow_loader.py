
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = PROJECT_ROOT / "data" / "business_workflows.xlsx"


def load_excel_sheets(file_path: Path = EXCEL_PATH) -> dict[str, list[dict[str, Any]]]:
    """Read non-empty Excel rows into dictionaries without pandas."""

    if not file_path.exists():
        raise FileNotFoundError(
            f"Workflow Excel not found: {file_path}\n"
            "Place the original workbook in the project's data folder."
        )

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    result: dict[str, list[dict[str, Any]]] = {}

    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                result[sheet.title] = []
                continue

            # Remove trailing empty header cells.
            headers = list(headers)
            while headers and headers[-1] is None:
                headers.pop()

            clean_headers = [
                str(value).strip() if value is not None else f"column_{i + 1}"
                for i, value in enumerate(headers)
            ]

            records = []

            for row_number, row in enumerate(rows, start=2):
                values = list(row[:len(clean_headers)])

                if not any(value is not None for value in values):
                    continue

                values.extend([None] * (len(clean_headers) - len(values)))

                record = dict(zip(clean_headers, values))
                record["_excel_row"] = row_number
                records.append(record)

            result[sheet.title] = records

    finally:
        workbook.close()

    return result


def preview_workflows() -> None:
    """Print the workbook structure for inspection."""

    sheets = load_excel_sheets()

    print(f"Workbook: {EXCEL_PATH.name}")
    print(f"Sheets found: {len(sheets)}")

    for sheet_name, records in sheets.items():
        print(f"\n{'=' * 60}")
        print(f"Sheet: {sheet_name}")
        print(f"Non-empty rows: {len(records)}")

        for record in records[:5]:
            print(record)


if __name__ == "__main__":
    try:
        preview_workflows()
    except (FileNotFoundError, OSError, ValueError) as exc:
        print(f"Excel loading error: {exc}")
