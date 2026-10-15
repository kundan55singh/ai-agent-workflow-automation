
import csv
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_INPUT_FILE_BYTES = 10 * 1024 * 1024

INVENTORY = [
    {"sku": "SKU-101", "name": "Wireless Mouse", "current_stock": 4, "reorder_threshold": 10},
    {"sku": "SKU-102", "name": "Keyboard", "current_stock": 25, "reorder_threshold": 10},
    {"sku": "SKU-103", "name": "USB-C Cable", "current_stock": 3, "reorder_threshold": 15},
]

PRICES = [
    {"sku": "SKU-101", "product_name": "Wireless Mouse", "internal_price": 499, "vendor_price": 520},
    {"sku": "SKU-102", "product_name": "Keyboard", "internal_price": 899, "vendor_price": 899},
    {"sku": "SKU-103", "product_name": "USB-C Cable", "internal_price": 199, "vendor_price": 250},
]

ORDERS = [
    {"order_id": "ORD-1001", "email": "customer@example.com", "status": "Shipped", "tracking_id": "TRK-1001"},
    {"order_id": "ORD-1002", "email": "buyer@example.com", "status": "Processing", "tracking_id": ""},
]

PRODUCTS = [
    {"sku": "SKU-101", "name": "Wireless Mouse", "category": "Electronics"},
    {"sku": "SKU-102", "name": "Mechanical Keyboard", "category": "Electronics"},
]

EMPLOYEES = [
    {"name": "Aarav", "skills": "inventory, reporting", "active_tasks": 2},
    {"name": "Diya", "skills": "marketing, content, seo", "active_tasks": 1},
    {"name": "Kabir", "skills": "data, reporting, analytics", "active_tasks": 3},
]

EXECUTION_HISTORY = [
    {"workflow_id": "WF001", "duration_seconds": 1.2, "status": "success"},
    {"workflow_id": "WF002", "duration_seconds": 2.4, "status": "success"},
    {"workflow_id": "WF003", "duration_seconds": 6.2, "status": "failed"},
]


# ==================================================
# FILE READING AND VALIDATION
# ==================================================

def read_tabular_file(file_path: str | Path) -> list[dict[str, Any]]:
    """Read CSV or Excel files safely and return rows as dictionaries."""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    if not path.is_file():
        raise ValueError(f"Input path is not a file: {path}")

    if path.suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError(f"Unsupported file type: {path.suffix}")

    if path.stat().st_size == 0:
        raise ValueError(f"Input file is empty: {path}")

    if path.stat().st_size > MAX_INPUT_FILE_BYTES:
        raise ValueError("Input file exceeds the 10 MB size limit")

    try:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as file:
                reader = csv.DictReader(file)

                if not reader.fieldnames:
                    raise ValueError("Input file has no header row")

                headers = [
                    str(header).strip() if header is not None else ""
                    for header in reader.fieldnames
                ]

                if any(not header for header in headers):
                    raise ValueError("Input file contains an empty column header")

                if len(headers) != len(set(headers)):
                    raise ValueError("Input file contains duplicate column headers")

                rows = []

                for row in reader:
                    if not row or all(
                        value is None or not str(value).strip()
                        for value in row.values()
                    ):
                        continue

                    rows.append({
                        str(key).strip(): (
                            value.strip() if isinstance(value, str) else value
                        )
                        for key, value in row.items()
                        if key is not None
                    })

        else:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )

            try:
                worksheet = workbook.active
                iterator = worksheet.iter_rows(values_only=True)
                header_row = next(iterator, None)

                if not header_row:
                    raise ValueError("Input file has no header row")

                headers = [
                    str(header).strip() if header is not None else ""
                    for header in header_row
                ]

                if any(not header for header in headers):
                    raise ValueError("Input file contains an empty column header")

                if len(headers) != len(set(headers)):
                    raise ValueError("Input file contains duplicate column headers")

                rows = []

                for values in iterator:
                    if not values or all(
                        value is None or not str(value).strip()
                        for value in values
                    ):
                        continue

                    rows.append(dict(zip(headers, values)))

            finally:
                workbook.close()

    except UnicodeDecodeError as exc:
        raise ValueError("CSV file must use UTF-8 encoding") from exc

    except InvalidFileException as exc:
        raise ValueError("Invalid or corrupted Excel file") from exc

    except (csv.Error, OSError) as exc:
        raise ValueError(f"Unable to read input file: {exc}") from exc

    if not rows:
        raise ValueError("Input file contains no data rows")

    return rows


# ==================================================
# COMMON HELPERS
# ==================================================

def _number(value: Any, field_name: str) -> float:
    """Convert input to a finite number with clear error handling."""
    try:
        if value is None or isinstance(value, bool) or str(value).strip() == "":
            raise ValueError

        number = float(value)

        if not math.isfinite(number):
            raise ValueError

        return number

    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(
            f"Invalid numeric value for '{field_name}': {value!r}"
        ) from exc


def _rows(
    context: dict,
    file_key: str,
    fallback: list[dict],
) -> list[dict]:
    """Read records from a file or use the workflow's demo data."""
    file_path = context.get(file_key)

    if file_path:
        return read_tabular_file(file_path)

    records = context.get(file_key.removesuffix("_file"), fallback)

    if not isinstance(records, list):
        raise ValueError(f"{file_key} data must be a list")

    if any(not isinstance(record, dict) for record in records):
        raise ValueError(f"Every record in {file_key} data must be a dictionary")

    return records


def _field(row: dict, *names: str, default=None):
    """Find a field using case-insensitive column names."""
    normalized = {
        str(key).strip().lower().replace(" ", "_"): value
        for key, value in row.items()
    }

    for name in names:
        key = name.strip().lower().replace(" ", "_")

        if key in normalized:
            return normalized[key]

    return default


def _success(data: Any, message: str = "Workflow completed successfully") -> dict:
    return {
        "status": "success",
        "message": message,
        "data": data,
    }


def _needs_input(message: str, required_fields: list[str]) -> dict:
    return {
        "status": "needs_input",
        "message": message,
        "required_fields": required_fields,
    }


# ==================================================
# WF001: INVENTORY RESTOCK CHECK
# ==================================================

def inventory_restock(context: dict) -> dict:
    records = _rows(context, "inventory_file", INVENTORY)
    restock = []

    for row in records:
        stock = _number(
            _field(row, "current_stock", "stock", "quantity"),
            "current_stock",
        )
        threshold = _number(
            _field(row, "reorder_threshold", "minimum_stock", "threshold"),
            "reorder_threshold",
        )

        if stock < 0 or threshold < 0:
            raise ValueError("Stock and reorder threshold cannot be negative")

        if stock <= threshold:
            restock.append({
                "sku": _field(row, "sku", "product_id", default="Unknown"),
                "product": _field(row, "name", "product_name", default="Unknown"),
                "current_stock": stock,
                "reorder_threshold": threshold,
                "recommended_order_quantity": max(
                    0, int(threshold * 2 - stock)
                ),
            })

    return {
        "status": "success",
        "message": f"Found {len(restock)} product(s) requiring a restock check",
        "count": len(restock),
        "data": restock,
    }


# ==================================================
# WF002: PRODUCT PRICE VALIDATION
# ==================================================

def validate_prices(context: dict) -> dict:
    records = _rows(context, "prices_file", PRICES)

    threshold = _number(
        context.get("threshold_percent", 10),
        "threshold_percent",
    )

    if threshold < 0:
        raise ValueError("threshold_percent cannot be negative")

    discrepancies = []

    for row in records:
        internal = _number(
            _field(row, "internal_price", "our_price"),
            "internal_price",
        )
        vendor = _number(
            _field(row, "vendor_price", "supplier_price"),
            "vendor_price",
        )

        if internal < 0 or vendor < 0:
            raise ValueError("Product prices cannot be negative")

        difference = abs(vendor - internal)
        percent = difference / internal * 100 if internal else (
            100.0 if vendor else 0.0
        )

        if percent > threshold:
            discrepancies.append({
                "sku": _field(row, "sku", default="Unknown"),
                "product": _field(row, "product_name", "name", default="Unknown"),
                "internal_price": internal,
                "vendor_price": vendor,
                "difference_percent": round(percent, 2),
            })

    return _success(discrepancies, "Price validation completed")


# ==================================================
# WF003: VENDOR FILE PROCESSING
# ==================================================

def process_vendor_file(context: dict) -> dict:
    file_path = context.get("vendor_file")

    if not file_path:
        return _needs_input(
            "Provide a vendor CSV or Excel file to process.",
            ["vendor_file"],
        )

    records = read_tabular_file(file_path)

    required_columns = {"sku", "price"}
    columns = {
        str(key).strip().lower().replace(" ", "_")
        for key in records[0]
    }

    missing = required_columns - columns

    if missing:
        return {
            "status": "error",
            "message": "Vendor file is missing required columns",
            "missing_columns": sorted(missing),
        }

    valid_records = []
    invalid_records = []

    for index, row in enumerate(records, start=2):
        try:
            sku = _field(row, "sku")
            price = _number(_field(row, "price"), "price")

            if not sku or price < 0:
                raise ValueError(
                    "SKU must be present and price cannot be negative"
                )

            valid_records.append({
                "sku": str(sku),
                "price": price,
            })

        except ValueError as exc:
            invalid_records.append({
                "row": index,
                "error": str(exc),
            })

    return _success({
        "total_rows": len(records),
        "valid_rows": valid_records,
        "invalid_rows": invalid_records,
    }, "Vendor file processing completed")


# ==================================================
# WF004: PRODUCT DESCRIPTION GENERATOR
# ==================================================

def generate_product_content(context: dict) -> dict:
    name = str(context.get("product_name", "")).strip()

    if not name:
        return _needs_input(
            "Provide the product name.",
            ["product_name"],
        )

    features = context.get("features", [])
    audience = str(
        context.get("target_audience", "everyday customers")
    ).strip()

    if isinstance(features, str):
        features = [
            item.strip()
            for item in features.split(",")
            if item.strip()
        ]

    if not isinstance(features, list):
        raise ValueError(
            "features must be a list or comma-separated string"
        )

    feature_text = ", ".join(str(item) for item in features)
    feature_text = feature_text or "practical design"

    description = (
        f"Discover the {name}, designed for {audience}. "
        f"Key features include {feature_text}. "
        "A practical choice for your everyday needs."
    )

    return _success({
        "product_name": name,
        "description": description,
        "seo_title": f"{name} | Features and Benefits",
        "seo_description": (
            f"Explore {name} and its key features: {feature_text}."
        ),
    }, "Product content generated")


# ==================================================
# WF005: CUSTOMER ORDER STATUS
# ==================================================

def get_order_status(context: dict) -> dict:
    order_id = str(context.get("order_id", "")).strip()
    email = str(context.get("email", "")).strip()

    if not order_id and not email:
        return _needs_input(
            "Provide an order ID or customer email.",
            ["order_id", "email"],
        )

    if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise ValueError("Invalid email address")

    matches = [
        order for order in ORDERS
        if (
            order_id
            and order["order_id"].lower() == order_id.lower()
        ) or (
            email
            and order["email"].lower() == email.lower()
        )
    ]

    if not matches:
        return {
            "status": "not_found",
            "message": "No matching order was found in the demo dataset",
        }

    return _success(matches, "Order lookup completed")


# ==================================================
# WF006: DUPLICATE PRODUCT DETECTION
# ==================================================

def detect_duplicates(context: dict) -> dict:
    records = _rows(context, "products_file", PRODUCTS)
    seen_skus = {}
    seen_names = {}
    duplicates = []

    for row in records:
        sku = str(
            _field(row, "sku", "product_id", default="")
        ).strip().lower()

        name = str(
            _field(row, "name", "product_name", default="")
        ).strip().lower()

        duplicate_of = None

        if sku and sku in seen_skus:
            duplicate_of = seen_skus[sku]
        elif name and name in seen_names:
            duplicate_of = seen_names[name]

        if duplicate_of:
            duplicates.append({
                "sku": _field(row, "sku", default=""),
                "name": _field(row, "name", "product_name", default=""),
                "duplicate_of": duplicate_of,
            })
        else:
            if sku:
                seen_skus[sku] = _field(row, "sku", default="")
            if name:
                seen_names[name] = (
                    _field(row, "sku", default="") or name
                )

    return _success(
        duplicates,
        f"Found {len(duplicates)} duplicate product(s)",
    )


# ==================================================
# WF007: MARKETING CAMPAIGN BRIEF
# ==================================================

def create_campaign_brief(context: dict) -> dict:
    goal = str(context.get("campaign_goal", "")).strip()

    if not goal:
        return _needs_input(
            "Provide a campaign goal.",
            ["campaign_goal"],
        )

    audience = str(
        context.get("target_audience", "general audience")
    ).strip()

    channel = str(
        context.get("channel", "social media")
    ).strip()

    brief = {
        "objective": goal,
        "target_audience": audience,
        "channel": channel,
        "key_message": f"Communicate a clear benefit to {audience}.",
        "call_to_action": "Learn more",
    }

    budget = context.get("budget")

    if budget not in (None, ""):
        amount = _number(budget, "budget")

        if amount < 0:
            raise ValueError("budget cannot be negative")

        brief["budget"] = amount

    return _success(brief, "Marketing campaign brief created")


# ==================================================
# WF008: SEO KEYWORD CLASSIFICATION
# ==================================================

def classify_keywords(context: dict) -> dict:
    keywords = context.get("keywords", [])

    if isinstance(keywords, str):
        keywords = [
            item.strip()
            for item in keywords.split(",")
            if item.strip()
        ]

    if not keywords and context.get("keyword_file"):
        records = read_tabular_file(context["keyword_file"])
        keywords = [
            str(_field(row, "keyword", "query", default="")).strip()
            for row in records
        ]
        keywords = [keyword for keyword in keywords if keyword]

    if not isinstance(keywords, list):
        raise ValueError(
            "keywords must be a list or comma-separated string"
        )

    if not keywords:
        return _needs_input(
            "Provide keywords to classify.",
            ["keywords"],
        )

    classified = []

    for keyword in keywords:
        text = str(keyword).strip()
        lower = text.lower()

        if not text:
            continue

        if any(word in lower for word in (
            "buy", "price", "discount", "order", "cheap"
        )):
            intent = "Transactional"

        elif any(word in lower for word in (
            "best", "versus", "vs", "review", "compare"
        )):
            intent = "Commercial"

        elif any(word in lower for word in (
            "how", "what", "why", "guide", "meaning"
        )):
            intent = "Informational"

        else:
            intent = "Navigational"

        classified.append({
            "keyword": text,
            "intent": intent,
        })

    return _success(
        classified,
        "Keyword classification completed",
    )


# ==================================================
# WF009: EMPLOYEE TASK ASSIGNMENT
# ==================================================

def assign_employee_task(context: dict) -> dict:
    task = str(context.get("task", "")).strip()

    if not task:
        return _needs_input(
            "Provide a task description.",
            ["task"],
        )

    task_lower = task.lower()

    skill_groups = {
        "marketing": ("marketing", "campaign", "content", "seo"),
        "inventory": ("inventory", "stock", "restock"),
        "data": ("data", "report", "analytics", "analysis"),
    }

    required_skills = [
        skill
        for skill, terms in skill_groups.items()
        if any(term in task_lower for term in terms)
    ]

    candidates = []

    for employee in EMPLOYEES:
        skills = employee["skills"].lower()

        if not required_skills or any(
            skill in skills for skill in required_skills
        ):
            candidates.append(employee)

    if not candidates:
        return {
            "status": "not_found",
            "message": "No employee matched the requested skills",
        }

    selected = min(
        candidates,
        key=lambda employee: employee["active_tasks"],
    )

    return _success({
        "task": task,
        "assigned_to": selected["name"],
        "matched_skills": required_skills,
        "current_active_tasks": selected["active_tasks"],
    }, "Task assignment recommendation generated")


# ==================================================
# WF010: WORKFLOW PERFORMANCE REPORT
# ==================================================

def workflow_performance(context: dict) -> dict:
    history = context.get(
        "execution_history",
        EXECUTION_HISTORY,
    )

    if not isinstance(history, list):
        raise ValueError("execution_history must be a list")

    slow_threshold = _number(
        context.get("slow_threshold_seconds", 5),
        "slow_threshold_seconds",
    )

    if slow_threshold < 0:
        raise ValueError(
            "slow_threshold_seconds cannot be negative"
        )

    slow_runs = []
    failures = []

    for record in history:
        if not isinstance(record, dict):
            raise ValueError(
                "Each execution history record must be a dictionary"
            )

        duration = _number(
            record.get("duration_seconds", 0),
            "duration_seconds",
        )

        if duration < 0:
            raise ValueError(
                "duration_seconds cannot be negative"
            )

        if duration > slow_threshold:
            slow_runs.append(record)

        if str(record.get("status", "")).lower() in {
            "failed", "error"
        }:
            failures.append(record)

    return _success({
        "total_executions": len(history),
        "slow_executions": slow_runs,
        "failed_executions": failures,
        "slow_threshold_seconds": slow_threshold,
    }, "Workflow performance report generated")


# ==================================================
# TOOL REGISTRY
# ==================================================

TOOL_REGISTRY = {
    "WF001": inventory_restock,
    "WF002": validate_prices,
    "WF003": process_vendor_file,
    "WF004": generate_product_content,
    "WF005": get_order_status,
    "WF006": detect_duplicates,
    "WF007": create_campaign_brief,
    "WF008": classify_keywords,
    "WF009": assign_employee_task,
    "WF010": workflow_performance,
}


def execute_tool(
    workflow_id: str,
    context: dict | None = None,
) -> dict:
    """Execute a registered workflow with input and result validation."""
    if not isinstance(workflow_id, str) or workflow_id not in TOOL_REGISTRY:
        raise ValueError(f"Unknown workflow: {workflow_id}")

    if context is None:
        context = {}

    if not isinstance(context, dict):
        raise ValueError("context must be a dictionary")

    result = TOOL_REGISTRY[workflow_id](context)

    if not isinstance(result, dict):
        raise TypeError(
            f"Workflow {workflow_id} returned a non-dictionary result"
        )

    return result
