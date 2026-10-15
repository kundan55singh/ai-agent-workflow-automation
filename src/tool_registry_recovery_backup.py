
import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_INPUT_FILE_BYTES = 10 * 1024 * 1024  # 10 MB


def read_tabular_file(path: str | Path) -> list[dict[str, Any]]:
    """Read CSV/XLSX safely and return clear validation errors."""

    if not str(path).strip():
        raise ValueError("Input file path cannot be empty.")

    file_path = Path(path).expanduser()

    if not file_path.exists():
        raise FileNotFoundError(
            f"Input file not found: {file_path}"
        )

    if not file_path.is_file():
        raise ValueError("Input path must point to a file.")

    if file_path.suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError(
            "Unsupported file type. Use CSV, XLSX or XLSM."
        )

    if file_path.stat().st_size > MAX_INPUT_FILE_BYTES:
        raise ValueError("Input file exceeds the 10 MB size limit.")

    if file_path.stat().st_size == 0:
        raise ValueError("Input file is empty.")

    if file_path.suffix.lower() == ".csv":
        try:
            with file_path.open(
                "r", encoding="utf-8-sig", newline=""
            ) as file:
                reader = csv.DictReader(file)

                if not reader.fieldnames:
                    raise ValueError(
                        "CSV must contain a header row."
                    )

                headers = [
                    str(header).strip()
                    if header is not None else ""
                    for header in reader.fieldnames
                ]

                if not any(headers):
                    raise ValueError(
                        "CSV headers cannot all be empty."
                    )

                return [
                    dict(zip(headers, row.values()))
                    for row in reader
                ]

        except UnicodeDecodeError as exc:
            raise ValueError(
                "CSV must use UTF-8 or UTF-8 BOM encoding."
            ) from exc
        except csv.Error as exc:
            raise ValueError(
                f"Could not parse CSV: {exc}"
            ) from exc

    workbook = None

    try:
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)

        if not headers or not any(
            value is not None and str(value).strip()
            for value in headers
        ):
            raise ValueError(
                "Excel worksheet must contain a header row."
            )

        names = [
            str(value).strip() if value is not None
            else f"column_{index}"
            for index, value in enumerate(headers)
        ]

        rows = []

        for row in iterator:
            if not any(value is not None for value in row):
                continue

            rows.append({
                names[index]: row[index]
                if index < len(row) else None
                for index in range(len(names))
            })

        return rows

    except (InvalidFileException, OSError) as exc:
        raise ValueError(
            "Excel file is invalid, corrupted or cannot be opened."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
